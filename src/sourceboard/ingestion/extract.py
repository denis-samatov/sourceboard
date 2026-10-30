"""Content extraction: turn source files into locatable pieces (page or time range)."""

from __future__ import annotations

import hashlib
import json
import logging
import re
import shutil
import subprocess
import threading
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from ..core.config import Settings

log = logging.getLogger(__name__)
MAX_DOCX_EXPANDED_BYTES = 128 * 1024 * 1024
MAX_IMAGE_PIXELS = 40_000_000


def validate_image(path: Path) -> None:
    """Inspect dimensions before OCR or an embedding model decodes pixels."""
    from PIL import Image

    with Image.open(path) as image:
        if image.width * image.height > MAX_IMAGE_PIXELS:
            raise ValueError("image pixel limit exceeded (40 million pixels)")


DOCUMENT_EXT = {".pdf", ".txt", ".md", ".markdown", ".docx"}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".opus"}
VIDEO_EXT = {".mp4", ".mov", ".m4v", ".mkv", ".webm", ".avi"}
SUBTITLE_EXT = {".srt", ".vtt"}


def kind_for(path: Path) -> str | None:
    ext = path.suffix.lower()
    if ext in DOCUMENT_EXT:
        return "document"
    if ext in IMAGE_EXT:
        return "image"
    if ext in AUDIO_EXT:
        return "audio"
    if ext in VIDEO_EXT:
        return "video"
    if ext in SUBTITLE_EXT:
        return "transcript"
    return None


@dataclass
class Piece:
    modality: str  # text | image | audio | video
    extractor: str
    text: str = ""
    page: int | None = None
    t_start: float | None = None
    t_end: float | None = None
    # File whose content is embedded for non-text modalities.
    media_path: Path | None = None
    # Persist media_path in the DB (frames, page renders) or treat it as temporary (audio clips).
    keep_media: bool = True


@dataclass
class Extraction:
    pieces: list[Piece] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    duration: float | None = None
    pages: int | None = None


def sha256_file(path: Path, chunk_size: int = 1 << 20) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        while block := fh.read(chunk_size):
            digest.update(block)
    return digest.hexdigest()


# --------------------------------------------------------------------------- text splitting

_SENTENCE_RE = re.compile(r"(?<=[.!?…])\s+(?=[\"«(\[]?[A-ZА-ЯЁ0-9])")


def split_text(text: str, max_chars: int) -> list[str]:
    """Split into paragraph-aligned chunks of at most ~max_chars characters."""
    paragraphs = [re.sub(r"[ \t]+", " ", p).strip() for p in re.split(r"\n\s*\n", text)]
    paragraphs = [p for p in paragraphs if p]
    units: list[str] = []
    for para in paragraphs:
        if len(para) <= max_chars:
            units.append(para)
            continue
        current = ""
        for sentence in _SENTENCE_RE.split(para):
            while len(sentence) > max_chars:  # pathological sentence: hard wrap
                cut = sentence.rfind(" ", 0, max_chars)
                cut = cut if cut > max_chars // 2 else max_chars
                if current:
                    units.append(current)
                    current = ""
                units.append(sentence[:cut].strip())
                sentence = sentence[cut:].strip()
            if current and len(current) + 1 + len(sentence) > max_chars:
                units.append(current)
                current = sentence
            else:
                current = f"{current} {sentence}".strip()
        if current:
            units.append(current)

    chunks: list[str] = []
    current = ""
    for unit in units:
        if current and len(current) + 2 + len(unit) > max_chars:
            chunks.append(current)
            current = unit
        else:
            current = f"{current}\n\n{unit}" if current else unit
    if current:
        chunks.append(current)
    return chunks


# --------------------------------------------------------------------------- subtitles

_TS_RE = re.compile(r"(?:(\d+):)?(\d{1,2}):(\d{2})[.,](\d{1,3})")


def parse_timestamp(value: str) -> float:
    match = _TS_RE.search(value)
    if not match:
        raise ValueError(f"Bad timestamp: {value!r}")
    hours, minutes, seconds, millis = match.groups()
    return (
        int(hours or 0) * 3600 + int(minutes) * 60 + int(seconds) + int(millis.ljust(3, "0")) / 1000
    )


@dataclass
class Cue:
    start: float
    end: float
    text: str


def parse_subtitles(text: str) -> list[Cue]:
    """Parse SRT or WebVTT content into cues."""
    cues: list[Cue] = []
    text = text.replace("\r\n", "\n").replace("\r", "\n").lstrip("﻿")
    for block in re.split(r"\n\s*\n", text):
        lines = [line for line in block.strip().split("\n") if line.strip()]
        timing_index = next((i for i, line in enumerate(lines) if "-->" in line), None)
        if timing_index is None:
            continue
        start_raw, end_raw = lines[timing_index].split("-->", 1)
        try:
            start = parse_timestamp(start_raw)
            end = parse_timestamp(end_raw.strip().split(" ")[0])
        except ValueError:
            continue
        body = " ".join(lines[timing_index + 1 :])
        body = re.sub(r"<[^>]+>", "", body).strip()  # VTT voice/style tags
        if body:
            cues.append(Cue(start, end, body))
    return cues


def window_cues(cues: list[Cue], window_s: float, max_chars: int, extractor: str) -> list[Piece]:
    """Group consecutive cues into windows aligned to cue boundaries."""
    pieces: list[Piece] = []
    group: list[Cue] = []

    def flush() -> None:
        if group:
            pieces.append(
                Piece(
                    modality="text",
                    extractor=extractor,
                    text=" ".join(c.text for c in group),
                    t_start=group[0].start,
                    t_end=group[-1].end,
                )
            )
            group.clear()

    for cue in cues:
        if group and (
            cue.end - group[0].start > window_s
            or sum(len(c.text) + 1 for c in group) + len(cue.text) > max_chars
        ):
            flush()
        group.append(cue)
    flush()
    return pieces


def cues_to_srt(cues: list[Cue]) -> str:
    def fmt(t: float) -> str:
        ms = round(t * 1000)
        return (
            f"{ms // 3_600_000:02d}:{ms // 60_000 % 60:02d}:{ms // 1000 % 60:02d},{ms % 1000:03d}"
        )

    return "\n".join(
        f"{i}\n{fmt(c.start)} --> {fmt(c.end)}\n{c.text}\n" for i, c in enumerate(cues, 1)
    )


# --------------------------------------------------------------------------- external tools


def tool(name: str) -> str | None:
    return shutil.which(name)


def _run(args: list[str], timeout: float = 600) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        args, capture_output=True, encoding="utf-8", errors="replace", timeout=timeout, check=True
    )


@lru_cache(maxsize=1)
def _tesseract_langs() -> frozenset[str]:
    exe = tool("tesseract")
    if not exe:
        return frozenset()
    try:
        out = subprocess.run(
            [exe, "--list-langs"],
            capture_output=True,
            encoding="utf-8",
            errors="replace",
            timeout=30,
        )
    except (OSError, subprocess.SubprocessError):
        return frozenset()
    return frozenset(line.strip() for line in out.stdout.splitlines()[1:] if line.strip())


def ocr_image(path: Path, langs: str) -> str | None:
    """OCR with the tesseract CLI. Returns None when tesseract is not installed."""
    exe = tool("tesseract")
    if not exe:
        return None
    available = _tesseract_langs()
    wanted = [lang for lang in langs.split("+") if lang in available] or ["eng"]
    try:
        # Resolved path: Leptonica rewrites paths that start with "/tmp" on macOS.
        out = _run([exe, str(path.resolve()), "stdout", "-l", "+".join(wanted)], timeout=180)
    except (OSError, subprocess.SubprocessError) as exc:
        log.warning("OCR failed for %s: %s", path, exc)
        return ""
    return re.sub(r"[ \t]+", " ", out.stdout).strip()


def media_duration(path: Path) -> float | None:
    exe = tool("ffprobe")
    if not exe:
        return None
    try:
        out = _run(
            [exe, "-v", "error", "-show_entries", "format=duration", "-of", "json", str(path)],
            timeout=60,
        )
        return float(json.loads(out.stdout)["format"]["duration"])
    except (OSError, subprocess.SubprocessError, KeyError, ValueError):
        return None


# --------------------------------------------------------------------------- ASR

_whisper_lock = threading.Lock()
_whisper_models: dict[str, object] = {}


def asr_available() -> bool:
    try:
        import faster_whisper  # noqa: F401
    except ImportError:
        return False
    return True


def transcribe(path: Path, model_size: str, cancel: threading.Event | None = None) -> list[Cue]:
    from faster_whisper import WhisperModel

    with _whisper_lock:
        model = _whisper_models.get(model_size)
        if model is None:
            model = WhisperModel(model_size, device="auto", compute_type="int8")
            _whisper_models[model_size] = model
        segments, _info = model.transcribe(str(path), vad_filter=True)  # type: ignore[attr-defined]
        cues = []
        for s in segments:  # a lazy generator: transcription happens while iterating
            if cancel is not None and cancel.is_set():
                raise IndexCancelled()
            if s.text.strip():
                cues.append(Cue(s.start, s.end, s.text.strip()))
        return cues


# --------------------------------------------------------------------------- extractors


class IndexCancelled(Exception):
    """Raised inside long extraction or embedding steps when the user stops indexing."""


class Extractor:
    def __init__(
        self,
        settings: Settings,
        modalities: frozenset[str],
        cancel: threading.Event | None = None,
    ) -> None:
        self.settings = settings
        self.modalities = modalities  # what the embedder can embed besides text
        self.cancel = cancel or threading.Event()

    def _check_cancel(self) -> None:
        if self.cancel.is_set():
            raise IndexCancelled()

    def extract(
        self, path: Path, kind: str, sha256: str, workdir: Path, sidecar: Path | None = None
    ) -> Extraction:
        if kind == "document":
            ext = path.suffix.lower()
            if ext == ".pdf":
                return self._pdf(path, sha256)
            if ext == ".docx":
                return self._docx(path)
            return self._plain(path)
        if kind == "transcript":
            return self._subtitles(path)
        if kind == "image":
            return self._image(path)
        if kind in {"audio", "video"}:
            return self._media(path, kind, sha256, workdir, sidecar)
        raise ValueError(f"Unsupported kind {kind}")

    def _text_pieces(self, text: str, extractor: str, page: int | None = None) -> list[Piece]:
        return [
            Piece(modality="text", extractor=extractor, text=chunk, page=page)
            for chunk in split_text(text, self.settings.text_chunk_chars)
        ]

    def _plain(self, path: Path) -> Extraction:
        text = path.read_text(encoding="utf-8", errors="replace")
        return Extraction(pieces=self._text_pieces(text, "plain"))

    def _docx(self, path: Path) -> Extraction:
        import docx

        with zipfile.ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_DOCX_EXPANDED_BYTES:
                raise ValueError("DOCX expanded size limit exceeded (128 MiB)")
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return Extraction(pieces=self._text_pieces("\n\n".join(parts), "docx"))

    def _pdf(self, path: Path, sha256: str) -> Extraction:
        from pypdf import PdfReader

        result = Extraction()
        reader = PdfReader(str(path))
        result.pages = len(reader.pages)
        empty_pages: list[int] = []
        for index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:  # pypdf raises a variety of errors on broken pages
                log.warning("pypdf failed on %s page %d: %s", path, index, exc)
                text = ""
            pieces = self._text_pieces(text, "pypdf", page=index)
            if not pieces:
                empty_pages.append(index)
            result.pieces.extend(pieces)

        need_ocr = bool(empty_pages) and self.settings.ocr_enabled and tool("tesseract")
        want_images = self.settings.render_pdf_pages and "image" in self.modalities
        if (need_ocr or want_images) and tool("pdftoppm"):
            out_dir = self.settings.media_dir / "pages" / sha256[:16]
            renders = self._render_pdf(path, out_dir, result.pages)
            if want_images:
                for page_no, image in renders.items():
                    result.pieces.append(
                        Piece(
                            modality="image",
                            extractor="page-render",
                            page=page_no,
                            media_path=image,
                        )
                    )
            if need_ocr:
                unreadable = 0
                for page_no in empty_pages:
                    self._check_cancel()
                    text = ""
                    if page_no in renders:
                        text = ocr_image(renders[page_no], self.settings.ocr_langs) or ""
                    pieces = self._text_pieces(text, "ocr", page=page_no)
                    unreadable += not pieces
                    result.pieces.extend(pieces)
                if unreadable:
                    result.notes.append(f"OCR не дал текста для {unreadable} стр.")
        elif empty_pages:
            result.notes.append(
                f"{len(empty_pages)} стр. без текстового слоя; для OCR нужны pdftoppm и tesseract"
            )
        return result

    def _render_pdf(self, path: Path, out_dir: Path, pages: int | None) -> dict[int, Path]:
        out_dir.mkdir(parents=True, exist_ok=True)
        existing = sorted(out_dir.glob("p-*.jpg"))
        if not existing or (pages and len(existing) < pages):
            try:
                _run(
                    [
                        "pdftoppm",
                        "-r",
                        "110",
                        "-jpeg",
                        "-jpegopt",
                        "quality=80",
                        str(path),
                        str(out_dir / "p"),
                    ],
                    timeout=1800,
                )
            except (OSError, subprocess.SubprocessError) as exc:
                log.warning("pdftoppm failed for %s: %s", path, exc)
                return {}
            existing = sorted(out_dir.glob("p-*.jpg"))
        return {int(p.stem.split("-")[-1]): p for p in existing}

    def _subtitles(self, path: Path) -> Extraction:
        cues = parse_subtitles(path.read_text(encoding="utf-8", errors="replace"))
        result = Extraction(duration=cues[-1].end if cues else None)
        result.pieces = window_cues(
            cues,
            self.settings.transcript_window_s,
            self.settings.text_chunk_chars,
            extractor=path.suffix.lower().lstrip("."),
        )
        return result

    def _image(self, path: Path) -> Extraction:
        validate_image(path)
        result = Extraction()
        text = ""
        if self.settings.ocr_enabled:
            ocr = ocr_image(path, self.settings.ocr_langs)
            if ocr is None:
                result.notes.append("OCR недоступен: установите tesseract")
            text = ocr or ""
        if "image" in self.modalities:
            result.pieces.append(
                Piece(modality="image", extractor="image", text=text, media_path=path)
            )
        elif text:
            # No image embeddings: index the OCR text so the image is still findable.
            result.pieces.extend(self._text_pieces(text, "ocr"))
        else:
            result.notes.append("изображение без текста и без визуального индекса")
        return result

    def _media(
        self, path: Path, kind: str, sha256: str, workdir: Path, sidecar: Path | None
    ) -> Extraction:
        result = Extraction(duration=media_duration(path))
        if not tool("ffmpeg"):
            result.notes.append("ffmpeg не найден: медиа индексируется только по субтитрам")

        cues = self._transcript_cues(path, sha256, sidecar, result)
        result.pieces.extend(
            window_cues(
                cues,
                self.settings.transcript_window_s,
                self.settings.text_chunk_chars,
                extractor="sidecar" if sidecar else "whisper",
            )
        )
        if result.duration is None and cues:
            result.duration = cues[-1].end

        if kind == "video" and "image" in self.modalities and tool("ffmpeg"):
            result.pieces.extend(self._video_frames(path, sha256, result.duration))
        if "audio" in self.modalities and tool("ffmpeg") and result.duration:
            result.pieces.extend(self._audio_windows(path, workdir, result.duration))
        return result

    def _transcript_cues(
        self, path: Path, sha256: str, sidecar: Path | None, result: Extraction
    ) -> list[Cue]:
        if sidecar is not None:
            return parse_subtitles(sidecar.read_text(encoding="utf-8", errors="replace"))
        cache = self.settings.media_dir / "transcripts" / f"{sha256}.srt"
        if cache.exists():
            return parse_subtitles(cache.read_text(encoding="utf-8"))
        if not (self.settings.asr_enabled and asr_available()):
            result.notes.append(
                "нет транскрипта: положите рядом .srt/.vtt с тем же именем или установите extra asr"
            )
            return []
        cues = transcribe(path, self.settings.asr_model, self.cancel)
        cache.parent.mkdir(parents=True, exist_ok=True)
        cache.write_text(cues_to_srt(cues), encoding="utf-8")
        return cues

    def _video_frames(self, path: Path, sha256: str, duration: float | None) -> list[Piece]:
        interval = self.settings.frame_interval_s
        out_dir = self.settings.media_dir / "frames" / f"{sha256[:16]}-{interval:g}"
        frames = sorted(out_dir.glob("f-*.jpg")) if out_dir.exists() else []
        if not frames:
            out_dir.mkdir(parents=True, exist_ok=True)
            try:
                # Frame n (1-based) is sampled at ~(n-1) * interval seconds.
                _run(
                    [
                        "ffmpeg",
                        "-nostdin",
                        "-loglevel",
                        "error",
                        "-i",
                        str(path),
                        "-vf",
                        f"fps=1/{interval:g},scale=640:-2",
                        "-q:v",
                        "4",
                        str(out_dir / "f-%06d.jpg"),
                    ],
                    timeout=3600,
                )
            except (OSError, subprocess.SubprocessError) as exc:
                log.warning("Frame extraction failed for %s: %s", path, exc)
                return []
            frames = sorted(out_dir.glob("f-*.jpg"))
        pieces = []
        for n, frame in enumerate(frames):
            start = n * interval
            end = start + interval
            if duration is not None:
                if start >= duration:
                    break
                end = min(end, duration)
            pieces.append(
                Piece(
                    modality="video", extractor="frame", t_start=start, t_end=end, media_path=frame
                )
            )
        return pieces

    def _audio_windows(self, path: Path, workdir: Path, duration: float) -> list[Piece]:
        window = self.settings.audio_window_s
        pieces = []
        start = 0.0
        index = 0
        while start < duration:
            self._check_cancel()
            length = min(window, duration - start)
            if length < 1.0 and pieces:
                break
            clip = workdir / f"a-{index:05d}.wav"
            try:
                _run(
                    [
                        "ffmpeg",
                        "-nostdin",
                        "-loglevel",
                        "error",
                        "-ss",
                        f"{start:.3f}",
                        "-t",
                        f"{length:.3f}",
                        "-i",
                        str(path),
                        "-vn",
                        "-ac",
                        "1",
                        "-ar",
                        "16000",
                        "-y",
                        str(clip),
                    ],
                    timeout=300,
                )
            except (OSError, subprocess.SubprocessError) as exc:
                log.warning("Audio clip extraction failed for %s at %.1fs: %s", path, start, exc)
                break
            pieces.append(
                Piece(
                    modality="audio",
                    extractor="audio-window",
                    t_start=start,
                    t_end=start + length,
                    media_path=clip,
                    keep_media=False,
                )
            )
            start += window
            index += 1
        return pieces
