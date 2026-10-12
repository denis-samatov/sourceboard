"""Generate a small fictional corpus for trying SourceBoard end to end.

Everything here is invented: the town, the plant, the people and the numbers.
Optional parts degrade gracefully: without a Cyrillic font the scanned PDF and slides are
skipped; without macOS `say` + ffmpeg the interviews are written as subtitles only.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
import wave
from pathlib import Path

from ..ingestion.extract import Cue, cues_to_srt

FONT_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "C:/Windows/Fonts/arial.ttf",
]

MINUTES = """# Протокол заседания совета депутатов посёлка Заречье

Дата: 12 декабря 2024 года. Присутствовали 9 депутатов из 11.

## Вопрос 1. О закрытии завода «Северный кабель»

Глава администрации сообщил, что собственник завода «Северный кабель» уведомил администрацию
о прекращении производства с 1 марта 2025 года. Под сокращение попадают 340 работников
основного производства.

Решили: просить областной центр занятости открыть программу переобучения для бывших
работников завода; рассмотреть сохранение льготного проезда для участников программы.

## Вопрос 2. О маршруте автобуса № 14

Перевозчик сообщил, что после остановки завода утренние рейсы маршрута № 14 до промзоны
станут убыточными. Предложено сократить число рейсов с 22 до 9 в сутки с апреля 2025 года.

Решили: согласовать сокращение рейсов при условии сохранения первого утреннего рейса в 6:10
и рейса до районной поликлиники.
"""

EDITOR_NOTE = """Заметка редактора: альтернативное объяснение падения пассажиропотока

Данные перевозчика показывают снижение пассажиропотока на маршруте № 14 уже в октябре
2024 года, то есть до объявления о закрытии завода. В октябре был закрыт на ремонт мост через
реку Сорму, и часть пассажиров перешла на электричку.

Нужно уточнить у перевозчика помесячные данные за 2024 год, прежде чем утверждать, что
сокращение рейсов вызвано только закрытием завода.
"""

TRANSPORT_ROWS = [
    ("Показатель", "Октябрь 2024", "Апрель 2025"),
    ("Рейсов в сутки, маршрут № 14", "22", "9"),
    ("Пассажиров в сутки", "1 850", "720"),
    ("Средняя наполняемость рейса", "84", "80"),
]

DECISION_PAGES = [
    [
        "АДМИНИСТРАЦИЯ ПОСЁЛКА ЗАРЕЧЬЕ",
        "ПОСТАНОВЛЕНИЕ № 47 от 20 января 2025 года",
        "",
        "О мерах поддержки работников завода",
        "«Северный кабель» в связи с прекращением",
        "производства",
        "",
        "1. Создать рабочую группу по трудоустройству",
        "340 высвобождаемых работников.",
        "2. Обратиться в центр занятости с просьбой",
        "организовать курсы переобучения сварщиков",
        "и операторов станков с ЧПУ.",
    ],
    [
        "3. Сохранить льготный проезд на маршруте № 14",
        "для участников программы переобучения",
        "до 31 декабря 2025 года.",
        "",
        "4. Контроль исполнения возложить",
        "на заместителя главы администрации.",
        "",
        "Глава администрации            В. Н. Карпов",
    ],
]

INTERVIEWS = {
    "ivanova_voditel": [
        "Я двенадцать лет вожу четырнадцатый маршрут.",
        "Раньше утром в автобус было не влезть, все ехали на завод к семи.",
        "После закрытия завода утренних рейсов оставили всего несколько.",
        "Сейчас люди ездят на работу в райцентр, и первый автобус в шесть десять всегда полный.",
        "Но честно скажу, пассажиров стало меньше ещё осенью, когда закрыли мост.",
    ],
    "petrov_master": [
        "Я работал мастером в кабельном цехе девятнадцать лет.",
        "Сокращения начались задолго до официального решения, уже летом двадцать "
        "четвёртого года людей отправляли в простой.",
        "Из моей бригады на переобучение пошли четверо, остальные уехали в город.",
        "Курсы сварщиков хорошие, но работы по этой специальности в посёлке почти нет.",
    ],
}

SLIDES = [
    ("Маршрут № 14", ["Рейсов в сутки:", "октябрь 2024 — 22", "апрель 2025 — 9"]),
    (
        "Программа переобучения",
        ["Записались: 112 человек", "Завершили курс: 64", "Трудоустроены: 31"],
    ),
]

MEETING_LINES = [
    "Добрый вечер, начинаем публичные слушания о транспорте.",
    "На слайде видно, что число рейсов четырнадцатого маршрута сократилось с двадцати "
    "двух до девяти.",
    "Следующий слайд про программу переобучения: записались сто двенадцать человек, "
    "трудоустроен тридцать один.",
]

EVAL_TASKS = [
    {
        "query": "когда завод прекращает производство",
        "expected": [{"source": "documents/protokol_soveta.md"}],
    },
    {
        "query": "сколько рейсов автобуса осталось после закрытия завода",
        "expected": [
            {"source": "documents/protokol_soveta.md"},
            {"source": "documents/spravka_transport.docx"},
            {"source": "video/slushaniya.mp4"},
        ],
    },
    {
        "query": "пассажиров стало меньше из-за ремонта моста",
        "expected": [
            {"source": "notes/zametka_redaktora.txt"},
            {"source": "interviews/ivanova_voditel.m4a"},
            {"source": "interviews/ivanova_voditel.srt"},
        ],
    },
    {
        "query": "увольнения начались раньше официального решения",
        "expected": [
            {"source": "interviews/petrov_master.m4a"},
            {"source": "interviews/petrov_master.srt"},
        ],
    },
    {
        "query": "льготный проезд для участников программы переобучения",
        "expected": [
            {"source": "documents/postanovlenie_47.pdf", "page": 2},
            {"source": "documents/protokol_soveta.md"},
        ],
    },
    {
        "query": "сколько человек трудоустроено после курсов",
        "expected": [
            {"source": "video/slushaniya.mp4"},
            {"source": "images/slide_pereobuchenie.png"},
        ],
    },
    {
        "query": "первый утренний рейс в 6:10",
        "expected": [
            {"source": "documents/protokol_soveta.md"},
            {"source": "interviews/ivanova_voditel.m4a"},
            {"source": "interviews/ivanova_voditel.srt"},
        ],
    },
    {
        "query": "курсы сварщиков и операторов ЧПУ",
        "expected": [
            {"source": "documents/postanovlenie_47.pdf", "page": 1},
            {"source": "interviews/petrov_master.m4a"},
            {"source": "interviews/petrov_master.srt"},
        ],
    },
]


def find_font() -> str | None:
    return next((f for f in FONT_CANDIDATES if Path(f).exists()), None)


def _image_with_lines(
    lines: list[str],
    size=(1240, 1754),
    title_size=44,
    body_size=36,
    font_path: str | None = None,
    title: str | None = None,
):
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    body = ImageFont.truetype(font_path, body_size) if font_path else ImageFont.load_default()
    y = 120
    if title:
        head = ImageFont.truetype(font_path, title_size) if font_path else body
        draw.text((100, y), title, fill="black", font=head)
        y += int(title_size * 2)
    for line in lines:
        draw.text((100, y), line, fill="black", font=body)
        y += int(body_size * 1.6)
    return image


def _write_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Справка перевозчика о маршруте № 14", level=1)
    document.add_paragraph(
        "После прекращения работы завода «Северный кабель» перевозчик сократил число рейсов "
        "маршрута № 14. Пассажиропоток снизился на 61 % по сравнению с октябрём 2024 года."
    )
    table = document.add_table(rows=0, cols=3)
    for row in TRANSPORT_ROWS:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = value
    document.add_paragraph("Сохранены первый утренний рейс в 6:10 и рейс до районной поликлиники.")
    document.save(str(path))


RATE = 22050
TOOL_TIMEOUT = 60  # `say` can hang when the speech service is unavailable


def _ffmpeg(*args: str) -> None:
    subprocess.run(
        ["ffmpeg", "-nostdin", "-loglevel", "error", "-y", *args],
        check=True,
        capture_output=True,
        timeout=600,
    )


def _speak(text: str, out_wav: Path) -> float:
    aiff = out_wav.with_suffix(".aiff")
    subprocess.run(
        ["say", "-v", "Milena", "-o", str(aiff), text],
        check=True,
        capture_output=True,
        timeout=TOOL_TIMEOUT,
    )
    _ffmpeg("-i", str(aiff), "-ac", "1", "-ar", str(RATE), "-c:a", "pcm_s16le", str(out_wav))
    with wave.open(str(out_wav)) as w:
        duration = w.getnframes() / w.getframerate()
    if duration < 0.2:  # e.g. speech service unavailable in a sandbox
        raise RuntimeError("say produced empty audio")
    return duration


def _write_wav(path: Path, frames: bytes) -> None:
    with wave.open(str(path), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(RATE)
        w.writeframes(frames)


def _narrate(lines: list[str], work: Path, gap: float = 0.8) -> tuple[Path, list[Cue]]:
    """Synthesize lines with exact cue timings and return (wav, cues)."""
    cues: list[Cue] = []
    t = 0.5
    frames = bytearray(b"\x00\x00" * int(RATE * t))
    for i, line in enumerate(lines):
        clip = work / f"line-{i}.wav"
        duration = _speak(line, clip)
        with wave.open(str(clip)) as w:
            frames += w.readframes(w.getnframes())
        cues.append(Cue(t, t + duration, line))
        t += duration + gap
        frames += b"\x00\x00" * int(RATE * gap)
    out = work / "narration.wav"
    _write_wav(out, bytes(frames))
    return out, cues


def _estimated_cues(lines: list[str], gap: float = 0.8) -> list[Cue]:
    """Plausible timings (~14 characters per second) when no speech synthesis is available."""
    cues, t = [], 1.0
    for line in lines:
        duration = max(2.0, len(line) / 14)
        cues.append(Cue(t, t + duration, line))
        t += duration + gap
    return cues


def _can_speak(work: Path) -> bool:
    """macOS `say` + ffmpeg present and actually producing sound."""
    if not (shutil.which("say") and shutil.which("ffmpeg")):
        return False
    try:
        _speak("Проверка.", work / "probe.wav")
    except (OSError, subprocess.SubprocessError, RuntimeError):
        return False
    return True


def _slideshow(target: Path, work: Path, cues: list[Cue], wav: Path | None) -> Path:
    """Two slides with a narration (or silent) track. Slide 2 starts with the third line."""
    total = cues[-1].end + 1.0
    if wav is None:
        wav = work / "silence.wav"
        _write_wav(wav, b"\x00\x00" * int(RATE * total))
    else:
        with wave.open(str(wav)) as w:
            total = w.getnframes() / w.getframerate()
    switch = cues[2].start - 0.4
    first, second = target / "images/slide_marshrut.png", target / "images/slide_pereobuchenie.png"
    concat = work / "slides.txt"
    concat.write_text(
        f"file '{first}'\nduration {switch:.3f}\n"
        f"file '{second}'\nduration {total - switch:.3f}\nfile '{second}'\n",
        encoding="utf-8",
    )
    video = target / "video/slushaniya.mp4"
    base = ["-f", "concat", "-safe", "0", "-i", str(concat), "-i", str(wav), "-r", "10"]
    base += ["-c:a", "aac", "-shortest"]
    try:
        _ffmpeg(*base, "-c:v", "libx264", "-pix_fmt", "yuv420p", str(video))
    except subprocess.CalledProcessError:  # ffmpeg built without libx264
        _ffmpeg(*base, "-c:v", "mpeg4", str(video))
    return video


def make_demo(target: Path) -> dict:
    target = target.expanduser().resolve()
    if target.exists() and any(target.iterdir()):
        raise FileExistsError(f"{target} уже существует и не пуст")
    for sub in ("documents", "interviews", "images", "notes", "video", "eval"):
        (target / sub).mkdir(parents=True, exist_ok=True)
    made: list[str] = []
    skipped: list[str] = []

    (target / "documents/protokol_soveta.md").write_text(MINUTES, encoding="utf-8")
    (target / "notes/zametka_redaktora.txt").write_text(EDITOR_NOTE, encoding="utf-8")
    _write_docx(target / "documents/spravka_transport.docx")
    made += [
        "documents/protokol_soveta.md",
        "notes/zametka_redaktora.txt",
        "documents/spravka_transport.docx",
    ]

    font = find_font()
    if font:
        pages = [_image_with_lines(lines, font_path=font) for lines in DECISION_PAGES]
        pdf = target / "documents/postanovlenie_47.pdf"
        pages[0].save(pdf, save_all=True, append_images=pages[1:], resolution=150)
        made.append("documents/postanovlenie_47.pdf (скан без текстового слоя)")
        for name, (title, lines) in zip(("marshrut", "pereobuchenie"), SLIDES, strict=True):
            slide = _image_with_lines(
                lines, size=(1280, 720), title_size=64, body_size=48, font_path=font, title=title
            )
            slide.save(target / f"images/slide_{name}.png")
            made.append(f"images/slide_{name}.png")
    else:
        skipped.append("скан PDF и слайды: не найден шрифт с кириллицей")

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        speech = _can_speak(work)
        if not speech:
            skipped.append("озвучка: нет рабочего macOS say + ffmpeg, таймкоды субтитров оценочные")
        for name, lines in INTERVIEWS.items():
            if speech:
                (work / name).mkdir()
                wav, cues = _narrate(lines, work / name)
                _ffmpeg(
                    "-i",
                    str(wav),
                    "-c:a",
                    "aac",
                    "-b:a",
                    "64k",
                    str(target / f"interviews/{name}.m4a"),
                )
                made.append(f"interviews/{name}.m4a")
            else:
                cues = _estimated_cues(lines)
            (target / f"interviews/{name}.srt").write_text(cues_to_srt(cues), encoding="utf-8")
            made.append(f"interviews/{name}.srt")

        if font and shutil.which("ffmpeg"):
            meeting = work / "meeting"
            meeting.mkdir()
            if speech:
                wav, cues = _narrate(MEETING_LINES, meeting)
            else:
                wav, cues = None, _estimated_cues(MEETING_LINES)
            _slideshow(target, meeting, cues, wav)
            (target / "video/slushaniya.srt").write_text(cues_to_srt(cues), encoding="utf-8")
            made += [
                "video/slushaniya.mp4" + ("" if speech else " (без звука)"),
                "video/slushaniya.srt",
            ]
        else:
            skipped.append("видео слушаний: нужны ffmpeg и шрифт с кириллицей")

    (target / "eval/tasks.jsonl").write_text(
        "\n".join(json.dumps(t, ensure_ascii=False) for t in EVAL_TASKS) + "\n", encoding="utf-8"
    )
    # No extension: keeps the note itself out of the index.
    (target / "README").write_text(
        "Демо-корпус SourceBoard. Все люди, организации, места и числа вымышлены.\n"
        "eval/tasks.jsonl — проверочные поисковые задания для `sourceboard eval`.\n",
        encoding="utf-8",
    )
    return {"path": str(target), "created": made, "skipped": skipped}
